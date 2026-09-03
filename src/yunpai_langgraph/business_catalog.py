from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .order_workbook import parse_order_workbook


SCHEMA_VERSION = "yunpai.business-catalog.v1"
IGNORED_NAMES = {".DS_Store"}
IGNORED_PREFIXES = ("._", "~$")
SUPPORTED_EXTENSIONS = {
    ".xlsx", ".xls", ".csv", ".tsv", ".json", ".pdf", ".docx", ".txt", ".md",
    ".dwg", ".et", ".zip", ".rar", ".7z", ".py", ".ps1",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_prefix(path: Path, limit: int = 65536) -> bytes:
    with path.open("rb") as stream:
        return stream.read(limit)


def _text_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (dict, list)):
        return "object" if isinstance(value, dict) else "array"
    return "string"


def classify_path(path: Path) -> tuple[str, str, float]:
    text = str(path).lower()
    name = path.name.lower()
    if any(token in text for token in ("订单", "备货", "定制单", "cg20", "po-20", "po_20")):
        return "order", "customer_or_stocking_order", 0.90
    if any(token in text for token in ("sop", "工艺", "作业指导")):
        return "sop", "production_sop", 0.93
    if any(token in text for token in ("bom", "物料清单")):
        return "bom", "engineering_bom", 0.93
    if any(token in text for token in ("承认书", "规格书", "datasheet", "数据手册")):
        return "engineering_document", "supplier_approval_or_spec", 0.86
    if any(token in text for token in ("采购", "po", "采购单", "采购订单")):
        return "procurement", "purchase_order_or_record", 0.82
    if any(token in text for token in ("库存", "入库", "出库", "盘点")):
        return "inventory", "inventory_record", 0.82
    if path.suffix.lower() in {".dwg"} or any(token in text for token in ("工程图", "图纸", "cad")):
        return "engineering_drawing", "cad_or_drawing", 0.88
    if path.suffix.lower() in {".xlsx", ".xls", ".csv", ".tsv"}:
        return "tabular", "unclassified_table", 0.45
    if path.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}:
        return "document", "unclassified_document", 0.40
    if path.suffix.lower() in {".zip", ".rar", ".7z"}:
        return "archive", "compressed_archive", 0.80
    return "other", "unclassified", 0.20


def _extract_xlsx(path: Path, kind: str) -> dict[str, Any]:
    from openpyxl import load_workbook

    result: dict[str, Any] = {"sheets": [], "active_sheet": None}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            nonempty = []
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 30), values_only=True):
                values = [str(value).strip()[:120] if value is not None else None for value in row[:40]]
                if any(value not in (None, "") for value in values):
                    nonempty.append(values)
            result["sheets"].append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "sample_rows": nonempty[:8],
            })
        result["active_sheet"] = workbook.active.title if workbook.active else None
        if kind == "order":
            try:
                result["order_document"] = parse_order_workbook(path.name, path.read_bytes())
            except Exception as exc:
                result["order_parse_error"] = str(exc)
    finally:
        workbook.close()
    return result


_BOM_HEADER_ALIASES = {
    "material_code": ("物料编码", "料号", "物料编号", "材料编码"),
    "material_name": ("材料名称", "原材料名称", "包材名称", "物料名称", "线材名称", "品名"),
    "specification": ("规格", "规格型号"),
    "quantity": ("用量", "数量", "用量/装箱数量"),
    "unit": ("单位",),
    "unit_price": ("单价", "含税单价", "不含税单价"),
    "cost": ("成本", "成本价格", "成本总价"),
    "supplier": ("供应商",),
}


def _looks_like_material_code(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    return bool(re.match(r"^(?:YA(?:\.[A-Z0-9]+)+|XC\d{3,}|[A-Z]{1,5}[._-][A-Z0-9._-]{2,})$", text, re.I))


def _header_key(value: Any) -> str | None:
    text = str(value or "").replace("\n", "").replace(" ", "").strip()
    if not text:
        return None
    for key, aliases in _BOM_HEADER_ALIASES.items():
        if any(alias.replace(" ", "") in text for alias in aliases):
            return key
    return None


def _extract_bom_xlsx(path: Path) -> dict[str, Any]:
    """Extract every populated BOM row from every sheet, retaining its locator."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    bom_lines: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[int, list[Any]]] = []
            max_nonempty_column = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [value for value in row]
                nonempty_columns = [index for index, value in enumerate(values, start=1) if value not in (None, "")]
                if not nonempty_columns:
                    continue
                max_nonempty_column = max(max_nonempty_column, max(nonempty_columns))
                rows.append((row_number, values))
            rows = [(number, values[:max_nonempty_column]) for number, values in rows]
            header_maps: list[tuple[int, dict[str, int]]] = []
            for number, values in rows[:20]:
                mapping = {key: index for index, value in enumerate(values) if (key := _header_key(value))}
                if "material_code" in mapping and ("material_name" in mapping or "quantity" in mapping):
                    header_maps.append((number, mapping))
            lines: list[dict[str, Any]] = []
            for row_number, values in rows:
                for header_row, mapping in header_maps:
                    if row_number <= header_row:
                        continue
                    code_index = mapping.get("material_code")
                    code = values[code_index] if code_index is not None and code_index < len(values) else None
                    if not _looks_like_material_code(code):
                        continue
                    line = {
                        "sheet_name": sheet.title,
                        "row_number": row_number,
                        "material_code": str(code).strip(),
                        "raw_cells": {str(index + 1): value for index, value in enumerate(values) if value not in (None, "")},
                    }
                    for key, index in mapping.items():
                        if index < len(values) and values[index] not in (None, ""):
                            line[key] = values[index]
                    lines.append(line)
                    break
            sheets.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "actual_max_column": max_nonempty_column,
                "nonempty_row_count": len(rows),
                "header_rows": [number for number, _ in header_maps],
                "bom_line_count": len(lines),
                "rows": [{"row_number": number, "values": values} for number, values in rows],
            })
            bom_lines.extend(lines)
    finally:
        workbook.close()
    return {"sheet_count": len(sheets), "sheets": sheets, "bom_lines": bom_lines, "bom_line_count": len(bom_lines)}


def _extract_delimited(path: Path) -> dict[str, Any]:
    encoding = "utf-8-sig"
    try:
        text = path.read_text(encoding=encoding, errors="replace")
    except OSError:
        return {"read_error": "unreadable"}
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if "\t" in sample else ","
    rows = list(csv.reader(text.splitlines()[:101], delimiter=delimiter))
    return {
        "delimiter": delimiter,
        "row_count_estimate": max(0, text.count("\n")),
        "headers": rows[0][:80] if rows else [],
        "sample_rows": [row[:80] for row in rows[1:6]],
    }


def _extract_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return {"parse_error": str(exc)}
    if isinstance(value, dict):
        return {"top_level_type": "object", "top_level_keys": list(value)[:100], "record_count": len(value.get("records", [])) if isinstance(value.get("records"), list) else None}
    if isinstance(value, list):
        return {"top_level_type": "array", "record_count": len(value), "first_record_keys": list(value[0])[:100] if value and isinstance(value[0], dict) else []}
    return {"top_level_type": _text_type(value)}


def _extract_text(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return {"text_chars": len(text), "text_preview": text[:2000]}
    if suffix == ".docx":
        try:
            from docx import Document
            document = Document(path)
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            tables = []
            for table in document.tables[:5]:
                tables.append([[cell.text.strip()[:200] for cell in row.cells[:30]] for row in table.rows[:20]])
            return {"paragraph_count": len(paragraphs), "text_preview": "\n".join(paragraphs)[:4000], "tables": tables}
        except Exception as exc:
            return {"parse_error": str(exc)}
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:5])
            return {"page_count": len(reader.pages), "text_preview": text[:4000], "text_chars_sample": len(text)}
        except Exception as exc:
            return {"parse_error": str(exc)}
    return {"binary_prefix_sha256": hashlib.sha256(_read_prefix(path)).hexdigest()}


def extract_file(path: Path, *, root: Path, deep_limit_bytes: int = 4_000_000, parse_xlsx: bool = False) -> dict[str, Any]:
    kind, subtype, classification_confidence = classify_path(path)
    suffix = path.suffix.lower()
    extraction: dict[str, Any]
    size = path.stat().st_size
    if suffix == ".xlsx" and (not parse_xlsx or size > deep_limit_bytes):
        extraction = {"extraction_skipped": "xlsx_deferred_to_m1_parser", "size_bytes": size}
    elif size > deep_limit_bytes and suffix in {".pdf", ".docx"}:
        extraction = {"extraction_skipped": "large_file", "size_bytes": size}
    elif suffix == ".xlsx":
        extraction = _extract_bom_xlsx(path) if kind == "bom" else _extract_xlsx(path, kind)
    elif suffix in {".csv", ".tsv"}:
        extraction = _extract_delimited(path)
    elif suffix == ".json":
        extraction = _extract_json(path)
    elif suffix in {".pdf", ".docx", ".txt", ".md"}:
        extraction = _extract_text(path)
    else:
        extraction = {"binary_prefix_sha256": hashlib.sha256(_read_prefix(path)).hexdigest()}
    document = extraction.get("order_document") if isinstance(extraction, dict) else None
    order = document if isinstance(document, dict) else {}
    lines = order.get("lines") if isinstance(order.get("lines"), list) else []
    field_observations = []
    for field in ("order_id", "order_date", "due_date", "supplier_name", "payment_terms", "delivery_address", "product_code", "quantity", "total_amount"):
        value = order.get(field)
        if value not in (None, ""):
            field_observations.append({"field_path": f"$.header.{field}", "raw_value": value, "normalized_value": value, "physical_type": _text_type(value), "semantic_type": field, "confidence": order.get("confidence", classification_confidence), "status": "candidate"})
    for line_index, line in enumerate(lines, start=1):
        for field, value in line.items():
            if value not in (None, ""):
                field_observations.append({"field_path": f"$.lines[{line_index - 1}].{field}", "raw_value": value, "normalized_value": value, "physical_type": _text_type(value), "semantic_type": field, "confidence": order.get("confidence", classification_confidence), "status": "candidate"})
    if kind == "bom" and isinstance(extraction, dict):
        document = {
            "schema_version": "m0.bom.v1",
            "document_type": "bom",
            "document_subtype": "engineering_bom",
            "confidence": 0.95 if extraction.get("bom_line_count") else 0.60,
            "review_status": "needs_review" if not extraction.get("bom_line_count") else "candidate",
            "sheet_count": extraction.get("sheet_count", 0),
            "bom_line_count": extraction.get("bom_line_count", 0),
        }
        for line_index, line in enumerate(extraction.get("bom_lines", [])):
            for field in ("material_code", "material_name", "specification", "quantity", "unit", "unit_price", "cost", "supplier"):
                if line.get(field) not in (None, ""):
                    field_observations.append({
                        "field_path": f"$.sheets[{line['sheet_name']!r}].rows[{line['row_number']}].{field}",
                        "raw_value": line[field], "normalized_value": line[field], "physical_type": _text_type(line[field]),
                        "semantic_type": field, "confidence": document["confidence"], "status": "candidate",
                    })
    return {
        "schema_version": SCHEMA_VERSION,
        "relative_path": str(path.relative_to(root)),
        "file_kind": kind,
        "document_subtype": subtype,
        "classification_confidence": classification_confidence,
        "extraction": extraction,
        "document": {
            "schema_version": document.get("schema_version") if isinstance(document, dict) and document.get("schema_version") else ("m1.document.v2" if order else None),
            "document_type": document.get("document_type") if isinstance(document, dict) and document.get("document_type") else (order.get("document_type") if order else kind),
            "document_subtype": document.get("document_subtype") if isinstance(document, dict) and document.get("document_subtype") else (order.get("document_subtype") if order else subtype),
            "order_id": order.get("order_id"),
            "product_code": order.get("product_code"),
            "confidence": document.get("confidence") if isinstance(document, dict) and document.get("confidence") is not None else order.get("confidence", classification_confidence),
            "review_status": document.get("review_status") if isinstance(document, dict) and document.get("review_status") else ("needs_review" if order else "unclassified"),
            **({"sheet_count": document.get("sheet_count"), "bom_line_count": document.get("bom_line_count")} if kind == "bom" and isinstance(document, dict) else {}),
        },
        "field_observations": field_observations,
    }


def iter_source_files(root: Path) -> Iterable[Path]:
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name in IGNORED_NAMES or any(path.name.startswith(prefix) for prefix in IGNORED_PREFIXES):
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        yield path


def init_catalog(db_path: str | Path) -> None:
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(path) as db:
        db.executescript(
            """
            PRAGMA foreign_keys = ON;
            CREATE TABLE IF NOT EXISTS ingest_batches (
                batch_id TEXT PRIMARY KEY,
                root_path TEXT NOT NULL,
                schema_version TEXT NOT NULL,
                started_at TEXT NOT NULL,
                finished_at TEXT,
                status TEXT NOT NULL,
                file_count INTEGER NOT NULL DEFAULT 0,
                summary_json TEXT NOT NULL DEFAULT '{}'
            );
            CREATE TABLE IF NOT EXISTS source_files (
                file_id TEXT PRIMARY KEY,
                batch_id TEXT NOT NULL REFERENCES ingest_batches(batch_id),
                absolute_path TEXT NOT NULL,
                relative_path TEXT NOT NULL,
                filename TEXT NOT NULL,
                extension TEXT NOT NULL,
                mime_type TEXT,
                size_bytes INTEGER NOT NULL,
                modified_at TEXT,
                sha256 TEXT NOT NULL,
                file_kind TEXT NOT NULL,
                document_subtype TEXT NOT NULL,
                classification_confidence REAL NOT NULL,
                status TEXT NOT NULL,
                metadata_json TEXT NOT NULL,
                UNIQUE(absolute_path, sha256)
            );
            CREATE TABLE IF NOT EXISTS document_candidates (
                document_id TEXT PRIMARY KEY,
                file_id TEXT NOT NULL REFERENCES source_files(file_id),
                schema_version TEXT NOT NULL,
                document_type TEXT NOT NULL,
                document_subtype TEXT NOT NULL,
                order_id TEXT,
                product_code TEXT,
                confidence REAL NOT NULL,
                review_status TEXT NOT NULL,
                payload_json TEXT NOT NULL,
                validation_issues_json TEXT NOT NULL DEFAULT '[]',
                UNIQUE(file_id)
            );
            CREATE TABLE IF NOT EXISTS field_observations (
                observation_id TEXT PRIMARY KEY,
                document_id TEXT NOT NULL REFERENCES document_candidates(document_id),
                field_path TEXT NOT NULL,
                raw_value_json TEXT NOT NULL,
                normalized_value_json TEXT,
                physical_type TEXT NOT NULL,
                semantic_type TEXT,
                unit TEXT,
                confidence REAL,
                status TEXT NOT NULL,
                evidence_json TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_source_files_kind ON source_files(file_kind);
            CREATE INDEX IF NOT EXISTS idx_document_candidates_order ON document_candidates(order_id);
            CREATE INDEX IF NOT EXISTS idx_field_observations_path ON field_observations(field_path);
            """
        )


def ingest_tree(root: str | Path, db_path: str | Path, *, batch_id: str | None = None, limit: int | None = None, parse_xlsx: bool = False, deep_limit_bytes: int = 4_000_000) -> dict[str, Any]:
    root_path = Path(root).expanduser().resolve()
    if not root_path.is_dir():
        raise ValueError(f"业务数据目录不存在: {root_path}")
    batch_id = batch_id or f"batch-{hashlib.sha256(f'{root_path}:{utc_now()}'.encode()).hexdigest()[:16]}"
    init_catalog(db_path)
    started_at = utc_now()
    paths = list(iter_source_files(root_path))
    if limit is not None:
        paths = paths[: max(0, limit)]
    counts: dict[str, int] = {}
    errors: list[dict[str, str]] = []
    with sqlite3.connect(db_path) as db:
        db.execute("PRAGMA foreign_keys = ON")
        db.execute("INSERT OR REPLACE INTO ingest_batches(batch_id, root_path, schema_version, started_at, status, file_count, summary_json) VALUES(?,?,?,?,?,?,?)", (batch_id, str(root_path), SCHEMA_VERSION, started_at, "running", 0, "{}"))
        for index, path in enumerate(paths):
            try:
                stat = path.stat()
                sha256 = _sha256(path)
                existing = db.execute("SELECT file_id, file_kind FROM source_files WHERE absolute_path=? AND sha256=?", (str(path), sha256)).fetchone()
                if existing:
                    counts[existing[1]] = counts.get(existing[1], 0) + 1
                    continue
                extracted = extract_file(path, root=root_path, parse_xlsx=parse_xlsx, deep_limit_bytes=deep_limit_bytes)
                source_key = hashlib.sha256(f"{path}\0{sha256}".encode()).hexdigest()
                file_id = f"file-{source_key[:24]}"
                document_id = f"doc-{source_key[:24]}"
                kind = extracted["file_kind"]
                counts[kind] = counts.get(kind, 0) + 1
                db.execute(
                    """INSERT INTO source_files(file_id,batch_id,absolute_path,relative_path,filename,extension,mime_type,size_bytes,modified_at,sha256,file_kind,document_subtype,classification_confidence,status,metadata_json)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(absolute_path,sha256) DO UPDATE SET batch_id=excluded.batch_id, metadata_json=excluded.metadata_json, status=excluded.status""",
                    (file_id, batch_id, str(path), extracted["relative_path"], path.name, path.suffix.lower(), mimetypes.guess_type(path.name)[0], stat.st_size, datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(), sha256, kind, extracted["document_subtype"], extracted["classification_confidence"], "identified", json_text({"schema_version": SCHEMA_VERSION, "extraction": extracted["extraction"]})),
                )
                db.execute("DELETE FROM field_observations WHERE document_id=?", (document_id,))
                document = extracted["document"]
                db.execute(
                    """INSERT INTO document_candidates(document_id,file_id,schema_version,document_type,document_subtype,order_id,product_code,confidence,review_status,payload_json,validation_issues_json)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(file_id) DO UPDATE SET document_id=excluded.document_id, payload_json=excluded.payload_json, order_id=excluded.order_id, product_code=excluded.product_code, confidence=excluded.confidence, review_status=excluded.review_status""",
                    (document_id, file_id, document.get("schema_version") or SCHEMA_VERSION, document.get("document_type") or kind, document.get("document_subtype") or extracted["document_subtype"], document.get("order_id"), document.get("product_code"), float(document.get("confidence") or 0), document.get("review_status") or "unclassified", json_text({"document": document, "extraction": extracted["extraction"]}), json_text([])),
                )
                for obs_index, observation in enumerate(extracted["field_observations"]):
                    obs_id = f"obs-{source_key[:16]}-{obs_index}"
                    db.execute(
                        "INSERT INTO field_observations(observation_id,document_id,field_path,raw_value_json,normalized_value_json,physical_type,semantic_type,unit,confidence,status,evidence_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                        (obs_id, document_id, observation["field_path"], json_text(observation["raw_value"]), json_text(observation["normalized_value"]), observation["physical_type"], observation.get("semantic_type"), observation.get("unit"), observation.get("confidence"), observation.get("status", "candidate"), json_text({"source_file": str(path), "sha256": sha256, "locator": observation["field_path"]})),
                    )
            except Exception as exc:
                errors.append({"path": str(path), "error": str(exc)})
                db.execute("INSERT OR REPLACE INTO source_files(file_id,batch_id,absolute_path,relative_path,filename,extension,mime_type,size_bytes,modified_at,sha256,file_kind,document_subtype,classification_confidence,status,metadata_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (f"error-{hashlib.sha256(str(path).encode()).hexdigest()[:24]}", batch_id, str(path), str(path.relative_to(root_path)), path.name, path.suffix.lower(), mimetypes.guess_type(path.name)[0], path.stat().st_size, None, "", "other", "unreadable", 0.0, "error", json_text({"error": str(exc)})))
            if (index + 1) % 25 == 0:
                db.commit()
        summary = {"counts_by_kind": counts, "errors": errors[:100], "error_count": len(errors), "root_path": str(root_path)}
        db.execute("UPDATE ingest_batches SET finished_at=?, status=?, file_count=?, summary_json=? WHERE batch_id=?", (utc_now(), "completed_with_errors" if errors else "completed", len(paths), json_text(summary), batch_id))
        db.commit()
    return {"batch_id": batch_id, "root_path": str(root_path), "file_count": len(paths), "counts_by_kind": counts, "error_count": len(errors), "errors": errors[:20], "db_path": str(Path(db_path).resolve())}


def catalog_summary(db_path: str | Path) -> dict[str, Any]:
    with sqlite3.connect(db_path) as db:
        return {
            "source_files": db.execute("SELECT count(*) FROM source_files").fetchone()[0],
            "documents": db.execute("SELECT count(*) FROM document_candidates").fetchone()[0],
            "field_observations": db.execute("SELECT count(*) FROM field_observations").fetchone()[0],
            "by_kind": db.execute("SELECT file_kind, count(*) FROM source_files GROUP BY file_kind ORDER BY file_kind").fetchall(),
            "by_review_status": db.execute("SELECT review_status, count(*) FROM document_candidates GROUP BY review_status ORDER BY review_status").fetchall(),
        }
