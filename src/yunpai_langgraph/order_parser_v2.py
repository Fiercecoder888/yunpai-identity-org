"""表头驱动的订单/业务表格解析：不依赖固定单元格坐标。

任务书 §3.2-3：订单解析不得依赖固定单元格坐标；建立表头别名、合并单元格、
隐藏行列、多 Sheet、日期/金额/数量归一化策略。
任务书 §3.2-4：数字列与包装/备注文本分离；`/`、空串、中文包装描述不得把整行
解析成 400 错误，也不能无证据静默转成 0。
任务书 §3.2-5/9：保存 parser_version、Sheet 名、行号/列号、原值、归一化值、
字段置信度与校验问题。

解析语义：
- 先找“明细表头行”（包含产品/数量等列别名的行）作为列映射；
- 表头行上方的“块事实”（订单号、交期、供应商、客户、日期等标签=值）作为表头字段；
- 表头行下方、非隐藏行、含产品编码或数量的行作为明细行；
- 不引用任何固定行号/列号（合并单元格经 MergedCellResolver 归一到左上角）。
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from openpyxl.utils import get_column_letter

PARSER_VERSION = "order.parser.v2"

# 语义字段 -> 表头别名（中英文与常见变体）。
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "order_id": ("订单号", "订单编号", "备货单号", "销售单号", "po号", "订单号码", "order no", "order_no", "po no", "po_number"),
    "order_date": ("下单日期", "订单日期", "制单日期", "order date", "order_date"),
    "due_date": ("交期", "交货日期", "交货期", "要求交期", "due date", "due_date", "delivery date", "需交期"),
    "customer": ("客户", "客户名称", "customer"),
    "supplier": ("供应商", "供应商名称", "supplier", "供方"),
    "product_code": ("型号", "料号", "产品型号", "产品编码", "物料编码", "编码", "part no", "part_no", "model", "产品编号", "物料号"),
    "product_name": ("名称", "品名", "产品名称", "物料名称", "描述", "name", "description", "材料名称", "货名"),
    "specification": ("规格", "规格型号", "spec", "size", "规格描述"),
    "quantity": ("数量", "采购数量", "订货数量", "订购数量", "需求数量", "qty", "quantity"),
    "unit": ("单位", "uom"),
    "unit_price": ("单价", "含税单价", "不含税单价", "price", "unit price"),
    "amount": ("金额", "总金额", "含税金额", "总价", "amount", "total", "合计金额"),
    "packaging": ("包装", "包装要求", "包装方式", "彩盒", "packing"),
    "remark": ("备注", "说明", "备注说明", "remark", "remark_note", "note", "备注栏"),
}

# 块事实（表头行上方 标签=值）；明细列别名不在此集合（避免把列当单值）。
BLOCK_FIELDS = ("order_id", "order_date", "due_date", "customer", "supplier")

_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日", "%Y.%m.%d", "%Y%m%d", "%m月%d日")
NULL_PLACEHOLDERS = {"", "/", "-", "--", "—", "n/a", "na", "无", "none", "null", "N/A", "NULL", "None"}

# 允许作为“块事实”标签的别名（排除与列别名冲突的宽泛词）。
_BLOCK_LABEL_ALIASES: dict[str, tuple[str, ...]] = {
    "order_id": ("订单号", "订单编号", "备货单号", "销售单号", "po号", "订单号码", "order no", "po no"),
    "order_date": ("下单日期", "订单日期", "制单日期", "order date"),
    "due_date": ("交期", "交货日期", "交货期", "要求交期", "due date"),
    "customer": ("客户", "客户名称"),
    "supplier": ("供应商", "供应商名称", "供方"),
}


class MergedCellResolver:
    """非 read_only workbook 的合并单元格：非左上角坐标归一到左上角。"""

    def __init__(self, sheet: Any) -> None:
        self._map: dict[tuple[int, int], tuple[int, int]] = {}
        for merged in getattr(sheet, "merged_cells", []):
            top_left = (merged.min_row, merged.min_col)
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    if (row, col) != top_left:
                        self._map[(row, col)] = top_left

    def coordinate(self, row: int, col: int) -> tuple[int, int]:
        return self._map.get((row, col), (row, col))


def _norm_label(text: str) -> str:
    return text.replace("\n", " ").replace(" ", "").replace("：", ":").strip().lower()


def header_field(value: Any, *, label_only: bool = False) -> str | None:
    text = _norm_label(str(value or ""))
    if not text:
        return None
    alias_sets = _BLOCK_LABEL_ALIASES if label_only else FIELD_ALIASES
    for field, aliases in alias_sets.items():
        if any(_norm_label(alias) == text for alias in aliases):
            return field
    if label_only:
        return None
    # 宽容子串匹配（如 “采购数量(个)”）。
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_norm = _norm_label(alias)
            if len(alias_norm) >= 2 and alias_norm in text:
                return field
    return None


def normalize_date(value: Any) -> tuple[str | None, str | None]:
    if value is None:
        return None, None
    if isinstance(value, datetime):
        return value.date().isoformat(), None
    if isinstance(value, date):
        return value.isoformat(), None
    if isinstance(value, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat(), None
        except (OverflowError, ValueError):
            return None, "INVALID_DATE"
    text = str(value).strip()
    if not text or text in NULL_PLACEHOLDERS:
        return None, None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat(), None
        except ValueError:
            continue
    return None, "INVALID_DATE"


def normalize_number(value: Any, *, issues: list[dict[str, Any]], field: str, sheet: str, row: int, col: int) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in NULL_PLACEHOLDERS:
        issues.append({
            "code": "NULL_PLACEHOLDER", "field": field, "sheet": sheet,
            "row": row, "column": col, "raw_value": text,
            "message": "占位符/空值未解析为数字，保留为 None（不静默转 0）",
        })
        return None
    try:
        return float(text.replace(",", "").replace("￥", "").replace("¥", "").strip())
    except ValueError:
        issues.append({
            "code": "INVALID_NUMBER", "field": field, "sheet": sheet,
            "row": row, "column": col, "raw_value": text,
        })
        return None


def _is_hidden_row(sheet: Any, row: int) -> bool:
    try:
        return bool(sheet.row_dimensions[row].hidden)
    except Exception:
        return False


def _is_hidden_col(sheet: Any, col: int) -> bool:
    try:
        return bool(sheet.column_dimensions[get_column_letter(col)].hidden)
    except Exception:
        return False


def _value_at(sheet: Any, resolver: MergedCellResolver, row: int, col: int) -> Any:
    resolved_row, resolved_col = resolver.coordinate(row, col)
    return sheet.cell(resolved_row, resolved_col).value


def _find_header(sheet: Any, resolver: MergedCellResolver, max_rows: int = 120, min_hits: int = 2) -> tuple[int | None, dict[str, int]]:
    """找明细表头行：行内至少有 2 个不同列别名命中。"""
    for row in range(1, min(sheet.max_row or 0, max_rows) + 1):
        if _is_hidden_row(sheet, row):
            continue
        mapping: dict[str, int] = {}
        for col in range(1, (sheet.max_column or 0) + 1):
            if _is_hidden_col(sheet, col):
                continue
            field = header_field(_value_at(sheet, resolver, row, col))
            if field and field not in mapping:
                mapping[field] = col
        if len(mapping) >= min_hits:
            return row, mapping
    return None, {}


def _scan_block_facts(sheet: Any, resolver: MergedCellResolver, stop_row: int) -> dict[str, Any]:
    """表头行上方扫描 标签=值 块事实（订单号/交期/供应商等）。

    取值顺序：标签同格“订单号：SO-001”取冒号后；否则取右邻格；否则取下邻格。
    """
    facts: dict[str, Any] = {}
    for row in range(1, max(1, stop_row)):
        for col in range(1, (sheet.max_column or 0) + 1):
            cell_value = _value_at(sheet, resolver, row, col)
            text = str(cell_value or "")
            label = header_field(cell_value, label_only=True)
            if label is None and ("：" in text or ":" in text):
                # “订单号：PO-BLOCK-1” 同格：先按冒号前缀匹配标签。
                prefix = re.split(r"[：:]", text, maxsplit=1)[0].strip()
                label = header_field(prefix, label_only=True)
            if label is None or label in facts:
                continue
            if "：" in text or ":" in text:
                candidate = re.split(r"[：:]", text, maxsplit=1)[1].strip()
                if candidate:
                    facts[label] = candidate
                    continue
            right = _value_at(sheet, resolver, row, col + 1)
            if right not in (None, ""):
                facts[label] = right
                continue
            below = _value_at(sheet, resolver, row + 1, col)
            if below not in (None, ""):
                facts[label] = below
    return facts


def parse_order_sheets(filename: str, raw: bytes, *, max_sheets: int = 16) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("XLSX 上传需要 openpyxl") from exc
    try:
        workbook = load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise ValueError(f"invalid XLSX file: {exc}") from exc

    try:
        sheet_docs: list[dict[str, Any]] = []
        all_lines: list[dict[str, Any]] = []
        header_facts: dict[str, Any] = {}
        validation_issues: list[dict[str, Any]] = []
        field_evidence: list[dict[str, Any]] = []
        sheets_with_headers = 0

        for sheet_index, sheet in enumerate(workbook.worksheets):
            if sheet_index >= max_sheets:
                break
            resolver = MergedCellResolver(sheet)
            header_row, mapping = _find_header(sheet, resolver)
            doc: dict[str, Any] = {"name": sheet.title, "header_row": header_row, "column_map": dict(mapping)}
            if header_row is None:
                doc["status"] = "no_header"
                sheet_docs.append(doc)
                continue
            doc["status"] = "header_found"
            sheet_docs.append(doc)
            sheets_with_headers += 1
            block_facts = _scan_block_facts(sheet, resolver, header_row)
            for field, value in block_facts.items():
                header_facts.setdefault(field, value)
            for row in range(header_row + 1, min(sheet.max_row or 0, 4000) + 1):
                if _is_hidden_row(sheet, row):
                    continue
                raw_cells: dict[str, Any] = {}
                for field, col in mapping.items():
                    raw_cells[field] = _value_at(sheet, resolver, row, col)
                # 列中出现的 order_id/due_date 等单值，在块事实缺失时回填为表头事实。
                for field in ("order_id", "order_date", "due_date", "customer", "supplier"):
                    if header_facts.get(field) is None and raw_cells.get(field) not in (None, ""):
                        header_facts[field] = raw_cells[field]
                has_product = raw_cells.get("product_code") not in (None, "")
                has_qty = normalize_number(
                    raw_cells.get("quantity"), issues=[], field="quantity", sheet=sheet.title, row=row, col=mapping.get("quantity", 0)
                ) is not None
                if not has_product and not has_qty:
                    continue
                line: dict[str, Any] = {"sheet": sheet.title, "row": row, "line_no": len(all_lines) + 1}
                for field in ("order_id", "product_code", "product_name", "specification", "unit", "packaging", "remark"):
                    line[field] = raw_cells.get(field)
                line["quantity"] = normalize_number(raw_cells.get("quantity"), issues=validation_issues, field="quantity", sheet=sheet.title, row=row, col=mapping.get("quantity", 0))
                line["unit_price"] = normalize_number(raw_cells.get("unit_price"), issues=validation_issues, field="unit_price", sheet=sheet.title, row=row, col=mapping.get("unit_price", 0))
                line["amount"] = normalize_number(raw_cells.get("amount"), issues=validation_issues, field="amount", sheet=sheet.title, row=row, col=mapping.get("amount", 0))
                if line.get("product_code") is None and raw_cells.get("product_name") not in (None, ""):
                    line["product_code"] = raw_cells.get("product_name")
                for field, col in mapping.items():
                    value = raw_cells.get(field)
                    if value in (None, ""):
                        continue
                    field_evidence.append({
                        "field_path": f"$.sheets[{sheet.title!r}].rows[{row}].{field}",
                        "sheet": sheet.title, "row": row, "column": col,
                        "raw_value": value, "semantic_type": field,
                        "parser_version": PARSER_VERSION,
                    })
                all_lines.append(line)

        for field, value in list(header_facts.items()):
            if field == "order_id":
                continue
            if field in {"order_date", "due_date"}:
                iso, issue = normalize_date(value)
                if issue:
                    validation_issues.append({"code": issue, "field": field, "sheet": "", "row": 0, "column": 0, "raw_value": value})
                    continue
                header_facts[field] = iso if iso else value
        return {
            "schema_version": "m1.document.v2",
            "document_type": "order",
            "document_subtype": "customer_or_stocking_order",
            "parser_version": PARSER_VERSION,
            "order_id": str(header_facts.get("order_id") or (all_lines[0].get("order_id") if all_lines and all_lines[0].get("order_id") else "") or ""),
            "order_date": str(header_facts.get("order_date") or ""),
            "due_date": str(header_facts.get("due_date") or ""),
            "supplier_name": str(header_facts.get("supplier") or header_facts.get("customer") or ""),
            "product_code": all_lines[0]["product_code"] if all_lines else None,
            "quantity": sum((line["quantity"] or 0) for line in all_lines),
            "lines": all_lines,
            "sheet_count": len(sheet_docs),
            "sheet_docs": sheet_docs,
            "field_evidence": field_evidence,
            "validation_issues": validation_issues,
            "confidence": 0.9 if all_lines and sheets_with_headers else 0.4,
            "source_filename": filename,
        }
    finally:
        workbook.close()
