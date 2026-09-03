from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any


def _value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def parse_order_workbook(filename: str, raw: bytes) -> dict[str, Any]:
    """Extract the stable order facts from the provided stocking-order workbook."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("XLSX upload requires openpyxl") from exc
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    cells = {cell.coordinate: _value(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None}

    def cell(address: str, default: Any = None) -> Any:
        value = cells.get(address, default)
        if isinstance(value, (int, float)) and address in {"X7"}:
            try:
                return from_excel(value, workbook.epoch).date().isoformat()
            except Exception:
                return value
        return value

    lines: list[dict[str, Any]] = []
    for row in range(10, sheet.max_row + 1):
        sequence = sheet.cell(row, 5).value
        model = _value(sheet.cell(row, 9).value)
        quantity = sheet.cell(row, 18).value
        if sequence in (None, "") or not model or quantity in (None, ""):
            continue
        lines.append({
            "line_id": f"{str(cells.get('P6') or 'order')}-{sequence}",
            "line_no": int(sequence) if str(sequence).isdigit() else sequence,
            "model": str(model),
            "product_code": str(model),
            "product_name": str(_value(sheet.cell(row, 10).value) or ""),
            "name_raw": str(_value(sheet.cell(row, 10).value) or ""),
            "specification": str(_value(sheet.cell(row, 14).value) or ""),
            "specification_raw": str(_value(sheet.cell(row, 14).value) or ""),
            "size": str(_value(sheet.cell(row, 17).value) or ""),
            "quantity": float(quantity),
            "unit": str(_value(sheet.cell(row, 19).value) or "PCS"),
            "unit_price": float(sheet.cell(row, 22).value or 0),
            "amount": float(sheet.cell(row, 23).value or 0),
            "unit_price_tax_included": float(sheet.cell(row, 22).value or 0),
            "amount_tax_included": float(sheet.cell(row, 23).value or 0),
            "dongguan_inventory": float(sheet.cell(row, 25).value or 0),
            "factory_inventory": float(sheet.cell(row, 26).value or 0),
            "in_transit": float(sheet.cell(row, 27).value or 0),
            "three_month_sales": float(sheet.cell(row, 28).value or 0),
            "packaging": str(_value(sheet.cell(row, 29).value) or ""),
            "pack_quantity": float(sheet.cell(row, 31).value or 0),
            "carton_count": float(sheet.cell(row, 32).value or 0),
            "remark": str(_value(sheet.cell(row, 34).value) or ""),
        })
    total_quantity = sum(line["quantity"] for line in lines)
    total_amount = sum(line["amount"] for line in lines)
    return {
        "order_id": str(cell("P6") or ""),
        "order_date": str(cell("X6") or ""),
        "due_date": str(cell("X7") or ""),
        "supplier_name": str(cell("G6") or ""),
        "payment_terms": str(cell("P7") or ""),
        "delivery_address": str(cell("D8") or ""),
        "product_code": lines[0]["product_code"] if lines else None,
        "quantity": total_quantity,
        "total_amount": total_amount,
        "lines": lines,
        "document_type": "order",
        "document_subtype": "stocking_order",
        "confidence": 0.98 if lines and cell("P6") and cell("X7") else 0.65,
        "source_filename": filename,
    }
