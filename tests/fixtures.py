"""本机阶段一：每格式脱敏 fixture 生成器（任务书 §五.1）。

覆盖：xlsx、xlsm(同族)、xls(OLE2 stub)、csv、tsv、json、pdf、docx、png、
jpg、zip、7z、rar(仅头)、坏文件、伪扩展名。全部为合成内容，不含真实生产
或个人资料。函数以 tmp 目录或 BytesIO 输出，供各测试复用。
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any


def xlsx_bytes(headers: list[str], rows: list[list[Any]], *, title: str = "Sheet1") -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def csv_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    import csv

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def tsv_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    output = io.StringIO()
    output.write("\t".join(headers) + "\n")
    for row in rows:
        output.write("\t".join(str(value) for value in row) + "\n")
    return output.getvalue().encode("utf-8")


def pdf_bytes(text: str = "PDF fixture page") -> bytes:
    return b"%PDF-1.7\n1 0 obj<</Type/Catalog>>endobj\n" + text.encode() + b"\n%%EOF\n"


def docx_bytes(text: str = "DOCX fixture paragraph") -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def png_bytes() -> bytes:
    # 最小合法 PNG（1x1 红色像素）。
    import base64

    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def jpg_bytes() -> bytes:
    return b"\xff\xd8\xff\xe0" + b"\x00" * 64


def zip_bytes(members: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return output.getvalue()


def nested_zip_bytes() -> bytes:
    inner = zip_bytes({"inner/inner.csv": b"a,b\n1,2\n"})
    return zip_bytes({"outer/order.xlsx": xlsx_bytes(["订单号", "型号", "数量"], [["SO-1", "P-1", 3]]), "outer/nested.zip": inner})


def seven_z_bytes() -> bytes:
    try:
        import py7zr
    except ImportError:  # pragma: no cover
        return b"7z\xbc\xaf\x27\x1c" + b"\x00" * 8
    output = io.BytesIO()
    with py7zr.SevenZipFile(output, "w") as archive:
        archive.writestr("a,b\n1,2\n", "member.csv")
    return output.getvalue()


def rar_stub_bytes() -> bytes:
    return b"Rar!\x1a\x07\x00" + b"\x00" * 16


def xls_stub_bytes() -> bytes:
    # OLE2 头 + 少量字节：sniff 识别为 xls；真实解析在缺有效 BIFF 流时报结构化错误。
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 96


def text_disguised_as_xlsx() -> bytes:
    return b"this is not a real xlsx, just text\nsecond line\n"


def corrupt_xlsx() -> bytes:
    return b"PK\x03\x04" + b"\x00" * 32


def write_all(tmp_path: Path) -> dict[str, Path]:
    """把所有 fixture 写入 tmp_path，返回 {name: path}。"""
    payload = {
        "sample.xlsx": xlsx_bytes(["物料编码", "材料名称", "用量"], [["M-1", "铜箔", 2]]),
        "sample.csv": csv_bytes(["订单号", "数量"], [["SO-1", 10]]),
        "sample.tsv": tsv_bytes(["型号", "数量"], [["P-1", 5]]),
        "sample.json": b'{"records":[{"kind":"order"}]}',
        "sample.pdf": pdf_bytes(),
        "sample.docx": docx_bytes(),
        "sample.png": png_bytes(),
        "sample.jpg": jpg_bytes(),
        "batch.zip": nested_zip_bytes(),
        "sample.7z": seven_z_bytes(),
        "sample.rar": rar_stub_bytes(),
        "legacy.xls": xls_stub_bytes(),
        "fake.xlsx": text_disguised_as_xlsx(),
        "broken.xlsx": corrupt_xlsx(),
    }
    written: dict[str, Path] = {}
    for name, content in payload.items():
        path = tmp_path / name
        path.write_bytes(content)
        written[name] = path
    return written
